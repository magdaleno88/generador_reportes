import importlib.util
import tempfile
import unittest
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
SPEC = importlib.util.spec_from_file_location("report_generator", ROOT / "generador aires.py")
GENERATOR = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(GENERATOR)


class ReportTests(unittest.TestCase):
    def test_folio_selection_expands_ranges_and_reports_duplicates(self):
        folios, warnings = GENERATOR.parsear_folios("100-102, 102; 105")
        self.assertEqual(folios, ["100", "101", "102", "105"])
        self.assertEqual(warnings, ["Folio repetido ignorado: 102"])

    def test_generated_word_report_uses_selected_equipment(self):
        with tempfile.TemporaryDirectory() as tmp:
            directory = Path(tmp)
            template = directory / "template.docx"
            source = Document()
            source.add_paragraph("Equipo {{FOLIO}}")
            table = source.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "Unidad"
            table.cell(0, 1).text = "{{UNIDAD}}"
            source.save(template)

            report = Path(GENERATOR.generar_reporte((
                0,
                {"folio": "DEMO/001", "unidad": "Unidad Ficticia", "servicio": "Preventivo"},
                [], template, directory,
            )))
            output = Document(report)
            self.assertEqual(report.name, "0001_demo_001_preventivo.docx")
            self.assertEqual(output.paragraphs[0].text, "Equipo DEMO/001")
            self.assertEqual(output.tables[0].cell(0, 1).text, "Unidad Ficticia")


if __name__ == "__main__":
    unittest.main()
