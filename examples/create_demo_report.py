"""Build a fictional Word report using the project's real report function."""

import importlib.util
import shutil
import tempfile
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt
from PIL import Image, ImageDraw


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = Path(__file__).resolve().parent / "demo-report.docx"


def load_generator():
    spec = importlib.util.spec_from_file_location("report_generator", ROOT / "generador aires.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def make_template(path, image_path):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)
    styles["Title"].font.name = "Arial"
    styles["Title"].font.size = Pt(20)

    doc.add_paragraph("Reporte de servicio de equipo", style="Title")
    doc.add_paragraph("Ejemplo de demostración con datos y una ilustración ficticios.")
    fields = [
        ("Folio", "{{FOLIO}}"),
        ("Unidad", "{{UNIDAD}}"),
        ("Región", "{{REGION}}"),
        ("Equipo", "{{EQUIPO}}"),
        ("Marca y modelo", "{{MARCA}} {{MODELO}}"),
        ("Serie", "{{SERIE}}"),
        ("Inventario", "{{INVENTARIO}}"),
        ("Servicio", "{{SERVICIO}}"),
    ]
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for label, placeholder in fields:
        row = table.add_row()
        row.cells[0].text = label
        row.cells[1].text = placeholder

    doc.add_paragraph("Evidencia de ejemplo", style="Heading 1")
    doc.add_paragraph("Ilustración creada para esta demostración; no representa un equipo ni una instalación reales.")
    doc.add_picture(str(image_path), width=Inches(5.9))
    doc.save(path)


def make_illustration(path):
    image = Image.new("RGB", (1200, 540), "#f4f1e9")
    draw = ImageDraw.Draw(image)
    draw.rectangle((150, 95, 1050, 435), fill="#e0e3e2", outline="#47555b", width=7)
    draw.rectangle((205, 150, 995, 350), fill="#f8faf9", outline="#7a898b", width=4)
    for y in range(180, 340, 25):
        draw.line((240, y, 850, y), fill="#9aabae", width=5)
    draw.ellipse((900, 220, 940, 260), fill="#b49350", outline="#47555b", width=3)
    draw.text((160, 465), "ILUSTRACION FICTICIA PARA DEMOSTRACION", fill="#39464b")
    image.save(path)


def main():
    generator = load_generator()
    with tempfile.TemporaryDirectory() as tmp:
        tmp_dir = Path(tmp)
        illustration = tmp_dir / "demo.png"
        template = tmp_dir / "template.docx"
        make_illustration(illustration)
        make_template(template, illustration)
        equipment = {
            "folio": "DEMO-001",
            "region": "Región Ejemplo",
            "unidad": "Centro de Salud Ficticio",
            "marca": "Marca Demo",
            "modelo": "MX-100",
            "serie": "SERIE-DEMO-0001",
            "inventario": "INV-DEMO-01",
            "equipo": "Minisplit de demostración",
            "servicio": "Preventivo",
            "tipo": "Equipo de climatización",
        }
        generated = generator.generar_reporte((0, equipment, [str(illustration)], template, tmp_dir))
        shutil.copyfile(generated, OUTPUT)
    print(f"Created {OUTPUT}")


if __name__ == "__main__":
    main()
