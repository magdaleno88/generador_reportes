import csv
import json
import multiprocessing
import os
import random
import re
import sqlite3
import threading
import tkinter as tk
import unicodedata
from concurrent.futures import ProcessPoolExecutor
from itertools import combinations, combinations_with_replacement
from math import comb
from pathlib import Path
from tkinter import filedialog, messagebox, ttk

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docxcompose.composer import Composer
from PIL import Image as PILImage

# ================= CONFIG =================

BASE_DIR = Path(__file__).resolve().parent
TEMPLATE_PATH = BASE_DIR / "template_2.docx"
OUTPUT_FOLDER = BASE_DIR / "output"
JSON_PATH = BASE_DIR / "base_datos_rf_ferems_CORREGIDA.json"
CSV_PATH = BASE_DIR / "base_datos_rf_ferems_CORREGIDA.csv"
OLD_JSON_PATH = BASE_DIR / "base_datos_rf_ferems.json"
OLD_CSV_PATH = BASE_DIR / "base_datos_rf_ferems.csv"
DB_PATH = BASE_DIR / "equipos.db"
IMAGENES_POR_REPORTE = 3
SERVICIOS = ("Preventivo", "Correctivo", "Ambos")
IMAGE_EXTENSIONS = (".jpg", ".jpeg", ".png")
MINISPLIT_SUBCARPETAS = ("completo", "compresor", "partes")

OUTPUT_FOLDER.mkdir(exist_ok=True)

# ================= CARPETAS DE FOTOS =================

RUTAS_FOTOS = {
    "MINISPLIT": BASE_DIR / "1-Minisplit",
    "UMA": BASE_DIR / "3-Uma",
    "PAQUETE": BASE_DIR / "2-Paquete",
    "CHILLER": BASE_DIR / "chiller raymundo 2026",
    "BOMBA": BASE_DIR / "bombas raymundo 2026",
    "TORRE": BASE_DIR / "torres 2026 raymundo",
    "EXTRACTOR": BASE_DIR / "5-extractor",
}

DB_COLUMNS = (
    "folio",
    "servicio",
    "region",
    "unidad",
    "marca",
    "modelo",
    "serie",
    "inventario",
    "equipo",
    "categoria",
    "tipo",
    "pagina_pdf",
)

# ================= UTILIDADES =================


def normalizar_texto(valor):
    texto = "" if valor is None else str(valor)
    texto = unicodedata.normalize("NFKD", texto)
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    return texto.casefold().strip()


def normalizar_clave(valor):
    texto = normalizar_texto(valor)
    texto = re.sub(r"[\.\-_]+", " ", texto)
    texto = re.sub(r"\s+", " ", texto).strip()
    return texto


def normalizar_categoria(valor):
    texto = normalizar_clave(valor)
    compacto = texto.replace(" ", "")

    if compacto.startswith("minisplit") or "mini split" in texto:
        return "MINISPLIT"
    if compacto.startswith("uma") or "unidad manejadora" in texto:
        return "UMA"
    if "unidad paquete" in texto or "unidad condensadora" in texto or texto == "paquete" or " paquete" in texto:
        return "PAQUETE"
    if "chiller" in texto or "enfriador" in texto:
        return "CHILLER"
    if "bomba" in texto:
        return "BOMBA"
    if "torre" in texto:
        return "TORRE"
    if "extractor" in texto:
        return "EXTRACTOR"
    return None


def sanitizar_nombre(valor, defecto="sin_nombre"):
    texto = normalizar_texto(valor)
    texto = re.sub(r'[<>:"/\\|?*\x00-\x1f]+', "_", texto)
    texto = re.sub(r"\s+", "_", texto)
    texto = re.sub(r"_+", "_", texto).strip("._ ")
    return texto or defecto


def servicio_archivo(servicio):
    return sanitizar_nombre(servicio.lower())


def escribir_log(log_path, mensaje):
    if log_path:
        with open(log_path, "a", encoding="utf-8") as f:
            f.write(str(mensaje) + "\n")


def obtener_imagenes(directorio):
    directorio = Path(directorio)
    if not directorio.exists():
        return []
    return [
        str(path)
        for path in directorio.iterdir()
        if path.is_file() and path.suffix.lower() in IMAGE_EXTENSIONS
    ]


def identificar_categoria(nombre):
    return normalizar_categoria(nombre)


def detectar_categoria_equipo(equipo):
    categoria_base = normalizar_categoria(equipo.get("categoria", ""))
    if categoria_base:
        return categoria_base
    return identificar_categoria(equipo.get("equipo", ""))


def categorias_estrictas():
    return ("MINISPLIT", "UMA", "PAQUETE", "CHILLER", "BOMBA", "TORRE", "EXTRACTOR")


# ================= BASE DE DATOS =================


def conectar_db():
    inicializar_db()
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def inicializar_db(force_rebuild=False):
    fuente_path, advertencia = resolver_fuente_datos()
    fuente_info = obtener_info_fuente(fuente_path)
    with sqlite3.connect(DB_PATH) as conn:
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS equipos (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                folio TEXT,
                servicio TEXT,
                region TEXT,
                unidad TEXT,
                marca TEXT,
                modelo TEXT,
                serie TEXT,
                inventario TEXT,
                equipo TEXT,
                categoria TEXT,
                tipo TEXT,
                pagina_pdf INTEGER,
                unidad_norm TEXT,
                folio_norm TEXT,
                inventario_norm TEXT,
                serie_norm TEXT
            )
            """
        )
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS metadata (
                clave TEXT PRIMARY KEY,
                valor TEXT
            )
            """
        )
        asegurar_columna(conn, "equipos", "categoria", "TEXT")
        crear_indices(conn)
        total = conn.execute("SELECT COUNT(*) FROM equipos").fetchone()[0]
        debe_reconstruir = force_rebuild or total == 0 or fuente_cambio(conn, fuente_info)
        if debe_reconstruir:
            registros = cargar_registros_fuente(fuente_path)
            importar_registros(conn, registros)
            guardar_metadata_fuente(conn, fuente_info)
            if advertencia:
                guardar_metadata(conn, "advertencia_fuente", advertencia)
            else:
                guardar_metadata(conn, "advertencia_fuente", "")
        elif advertencia:
            guardar_metadata(conn, "advertencia_fuente", advertencia)
        return resumen_base(conn)


def asegurar_columna(conn, tabla, columna, tipo):
    columnas = [row[1] for row in conn.execute(f"PRAGMA table_info({tabla})").fetchall()]
    if columna not in columnas:
        conn.execute(f"ALTER TABLE {tabla} ADD COLUMN {columna} {tipo}")


def crear_indices(conn):
    indices = {
        "idx_equipos_unidad": "unidad",
        "idx_equipos_servicio": "servicio",
        "idx_equipos_folio": "folio",
        "idx_equipos_inventario": "inventario",
        "idx_equipos_serie": "serie",
        "idx_equipos_unidad_norm": "unidad_norm",
        "idx_equipos_categoria": "categoria",
    }
    for nombre, columna in indices.items():
        conn.execute(f"CREATE INDEX IF NOT EXISTS {nombre} ON equipos ({columna})")


def resolver_fuente_datos():
    if JSON_PATH.exists():
        return JSON_PATH, ""
    if CSV_PATH.exists():
        return CSV_PATH, "Advertencia: no existe base_datos_rf_ferems_CORREGIDA.json; se usara base_datos_rf_ferems_CORREGIDA.csv."
    if OLD_JSON_PATH.exists():
        return OLD_JSON_PATH, "Advertencia: no existen archivos corregidos; se usara base_datos_rf_ferems.json como ultimo recurso."
    if OLD_CSV_PATH.exists():
        return OLD_CSV_PATH, "Advertencia: no existen archivos corregidos; se usara base_datos_rf_ferems.csv como ultimo recurso."
    raise FileNotFoundError(
        "No se encontro base_datos_rf_ferems_CORREGIDA.json ni base_datos_rf_ferems_CORREGIDA.csv"
    )


def obtener_info_fuente(fuente_path):
    stat = fuente_path.stat()
    return {
        "source_file": fuente_path.name,
        "source_path": str(fuente_path.resolve()),
        "source_mtime": str(stat.st_mtime_ns),
        "source_size": str(stat.st_size),
    }


def guardar_metadata(conn, clave, valor):
    conn.execute(
        "INSERT OR REPLACE INTO metadata (clave, valor) VALUES (?, ?)",
        (clave, "" if valor is None else str(valor)),
    )
    conn.commit()


def leer_metadata(conn, clave, defecto=""):
    row = conn.execute("SELECT valor FROM metadata WHERE clave = ?", (clave,)).fetchone()
    return row[0] if row else defecto


def fuente_cambio(conn, fuente_info):
    for clave, valor in fuente_info.items():
        if leer_metadata(conn, clave) != valor:
            return True
    return False


def guardar_metadata_fuente(conn, fuente_info):
    for clave, valor in fuente_info.items():
        conn.execute(
            "INSERT OR REPLACE INTO metadata (clave, valor) VALUES (?, ?)",
            (clave, valor),
        )
    conn.commit()


def resumen_base(conn):
    total = conn.execute("SELECT COUNT(*) FROM equipos").fetchone()[0]
    total_unidades = conn.execute("SELECT COUNT(DISTINCT unidad) FROM equipos").fetchone()[0]
    return {
        "source_file": leer_metadata(conn, "source_file"),
        "total_registros": total,
        "total_unidades": total_unidades,
        "advertencia": leer_metadata(conn, "advertencia_fuente"),
    }


def cargar_registros_fuente(fuente_path):
    if fuente_path.suffix.lower() == ".json":
        try:
            with open(fuente_path, "r", encoding="utf-8-sig") as f:
                registros = json.load(f)
            if isinstance(registros, list):
                return registros
        except Exception:
            pass

    if fuente_path.suffix.lower() == ".csv":
        with open(fuente_path, "r", encoding="utf-8-sig", newline="") as f:
            return list(csv.DictReader(f))

    raise ValueError(f"Fuente de datos no soportada: {fuente_path}")


def importar_registros(conn, registros):
    conn.execute("DELETE FROM equipos")
    filas = []
    for registro in registros:
        fila = {col: registro.get(col, "") for col in DB_COLUMNS}
        fila["categoria"] = normalizar_categoria(fila.get("categoria", "")) or identificar_categoria(fila.get("equipo", "")) or ""
        try:
            fila["pagina_pdf"] = int(fila.get("pagina_pdf") or 0)
        except (TypeError, ValueError):
            fila["pagina_pdf"] = 0
        fila["unidad_norm"] = normalizar_texto(fila["unidad"])
        fila["folio_norm"] = normalizar_texto(fila["folio"])
        fila["inventario_norm"] = normalizar_texto(fila["inventario"])
        fila["serie_norm"] = normalizar_texto(fila["serie"])
        filas.append(fila)

    conn.executemany(
        """
        INSERT INTO equipos (
            folio, servicio, region, unidad, marca, modelo, serie, inventario,
            equipo, categoria, tipo, pagina_pdf, unidad_norm, folio_norm, inventario_norm, serie_norm
        ) VALUES (
            :folio, :servicio, :region, :unidad, :marca, :modelo, :serie, :inventario,
            :equipo, :categoria, :tipo, :pagina_pdf, :unidad_norm, :folio_norm, :inventario_norm, :serie_norm
        )
        """,
        filas,
    )
    conn.commit()


def filas_a_dicts(rows):
    return [dict(row) for row in rows]


def buscar_unidades(texto_busqueda):
    termino = normalizar_texto(texto_busqueda)
    with conectar_db() as conn:
        if not termino:
            rows = conn.execute(
                "SELECT unidad, COUNT(*) total FROM equipos GROUP BY unidad ORDER BY unidad LIMIT 100"
            ).fetchall()
        else:
            rows = conn.execute(
                """
                SELECT unidad, COUNT(*) total
                FROM equipos
                WHERE unidad_norm LIKE ?
                GROUP BY unidad
                ORDER BY unidad
                LIMIT 100
                """,
                (f"%{termino}%",),
            ).fetchall()
    return [row["unidad"] for row in rows]


def obtener_equipos(unidad, servicio):
    params = [unidad]
    where = ["unidad = ?"]
    if servicio in ("Preventivo", "Correctivo"):
        where.append("servicio = ?")
        params.append(servicio)

    with conectar_db() as conn:
        rows = conn.execute(
            f"SELECT {', '.join(DB_COLUMNS)} FROM equipos WHERE {' AND '.join(where)} ORDER BY folio, servicio",
            params,
        ).fetchall()
    return filas_a_dicts(rows)


def contar_servicios(equipos):
    preventivos = sum(1 for e in equipos if normalizar_texto(e.get("servicio")) == "preventivo")
    correctivos = sum(1 for e in equipos if normalizar_texto(e.get("servicio")) == "correctivo")
    return preventivos, correctivos


# ================= GESTOR DE FOTOS =================


class GestorFotos:
    def __init__(self, rutas_fotos=None, log=None):
        self.pools = {}
        self.minisplit_pools = {}
        self.rutas_fotos = {}
        self.log = log or (lambda mensaje: None)
        rutas_recibidas = RUTAS_FOTOS if rutas_fotos is None else rutas_fotos
        for cat, ruta in rutas_recibidas.items():
            categoria = normalizar_categoria(cat)
            if not categoria:
                self.log(f"Advertencia: se ignora carpeta con categoria desconocida: {cat} -> {ruta}")
                continue
            if categoria not in categorias_estrictas():
                self.log(f"Advertencia: categoria no permitida para fotos: {cat} -> {ruta}")
                continue
            self.rutas_fotos[categoria] = Path(ruta)
        for categoria in categorias_estrictas():
            if categoria not in self.rutas_fotos:
                self.log(f"Advertencia: no hay ruta configurada para categoria {categoria}")
        for categoria, ruta in self.rutas_fotos.items():
            if categoria == "MINISPLIT":
                self._cargar_minisplit(ruta)
                continue
            fotos = obtener_imagenes(ruta)
            if not ruta.exists():
                self.log(f"Carpeta no encontrada para categoria {categoria}: {ruta}")
            elif len(fotos) == 0:
                self.log(f"Carpeta sin imagenes validas para categoria {categoria}: {ruta}")
            random.shuffle(fotos)
            self.pools[categoria] = {
                "todas": fotos.copy(),
                "por_cantidad": {},
            }

    def _cargar_minisplit(self, ruta):
        if not ruta.exists():
            self.log(f"Carpeta no encontrada para categoria MINISPLIT: {ruta}")
        for subcarpeta in MINISPLIT_SUBCARPETAS:
            subruta = ruta / subcarpeta
            fotos = obtener_imagenes(subruta)
            if not subruta.exists():
                self.log(f"Carpeta no encontrada para MINISPLIT/{subcarpeta}: {subruta}")
            elif len(fotos) == 0:
                self.log(f"Carpeta sin imagenes validas para MINISPLIT/{subcarpeta}: {subruta}")
            random.shuffle(fotos)
            self.minisplit_pools[subcarpeta] = {
                "ruta": subruta,
                "todas": fotos.copy(),
                "disponibles": fotos.copy(),
            }

    def _obtener_una_minisplit(self, subcarpeta):
        pool = self.minisplit_pools.get(subcarpeta)
        if not pool or not pool["todas"]:
            return None
        if not pool["disponibles"]:
            pool["disponibles"] = pool["todas"].copy()
            random.shuffle(pool["disponibles"])
        return pool["disponibles"].pop()

    @staticmethod
    def _total_combinaciones(total_fotos, cantidad):
        if total_fotos == 0:
            return 0
        if total_fotos >= cantidad:
            return comb(total_fotos, cantidad)
        return comb(total_fotos + cantidad - 1, cantidad)

    @staticmethod
    def _crear_cache_indices(total_fotos, cantidad):
        if total_fotos >= cantidad:
            indices = list(combinations(range(total_fotos), cantidad))
        else:
            indices = list(combinations_with_replacement(range(total_fotos), cantidad))
        random.shuffle(indices)
        return indices

    def _obtener_estado(self, pool, cantidad):
        if cantidad not in pool["por_cantidad"]:
            pool["por_cantidad"][cantidad] = {
                "usadas": set(),
                "cache": None,
                "cursor": 0,
                "total": self._total_combinaciones(len(pool["todas"]), cantidad),
            }
        return pool["por_cantidad"][cantidad]

    def _obtener_indices(self, pool, cantidad):
        estado = self._obtener_estado(pool, cantidad)
        total_fotos = len(pool["todas"])
        if estado["total"] == 0:
            return None

        if len(estado["usadas"]) >= estado["total"]:
            estado["usadas"].clear()
            estado["cache"] = None
            estado["cursor"] = 0

        if estado["total"] <= 200000:
            if estado["cache"] is None:
                estado["cache"] = self._crear_cache_indices(total_fotos, cantidad)
                estado["cursor"] = 0

            while estado["cursor"] < len(estado["cache"]):
                indices = estado["cache"][estado["cursor"]]
                estado["cursor"] += 1
                if indices not in estado["usadas"]:
                    estado["usadas"].add(indices)
                    return indices
        else:
            for _ in range(5000):
                if total_fotos >= cantidad:
                    indices = tuple(sorted(random.sample(range(total_fotos), cantidad)))
                else:
                    indices = tuple(sorted(random.choices(range(total_fotos), k=cantidad)))
                if indices not in estado["usadas"]:
                    estado["usadas"].add(indices)
                    return indices

            for indices in combinations(range(total_fotos), cantidad):
                if indices not in estado["usadas"]:
                    estado["usadas"].add(indices)
                    return indices
        return None

    def obtener(self, categoria, cantidad=IMAGENES_POR_REPORTE):
        categoria = normalizar_categoria(categoria)
        if categoria == "MINISPLIT":
            return self.obtener_minisplit()
        if categoria not in self.pools:
            self.log(f"Advertencia: categoria no encontrada para fotos: {categoria}")
            return []
        pool = self.pools[categoria]
        if len(pool["todas"]) == 0:
            self.log(f"Advertencia: categoria sin fotos: {categoria}")
            return []

        indices = self._obtener_indices(pool, cantidad)
        if indices is None:
            return []
        seleccion = [pool["todas"][idx] for idx in indices]
        random.shuffle(seleccion)
        return seleccion

    def obtener_minisplit(self):
        seleccion = []
        for subcarpeta in MINISPLIT_SUBCARPETAS:
            foto = self._obtener_una_minisplit(subcarpeta)
            if foto:
                seleccion.append(foto)
            else:
                pool = self.minisplit_pools.get(subcarpeta)
                ruta = pool["ruta"] if pool else self.rutas_fotos.get("MINISPLIT", Path("SIN_RUTA")) / subcarpeta
                self.log(
                    f"Advertencia: no hay imagen disponible para MINISPLIT/{subcarpeta}: {ruta}"
                )
        return seleccion

    def obtener_para_equipo(self, equipo, cantidad=IMAGENES_POR_REPORTE):
        categoria = detectar_categoria_equipo(equipo)
        carpeta = self.rutas_fotos.get(categoria)
        detalle = {
            "folio": equipo.get("folio", ""),
            "equipo": equipo.get("equipo", ""),
            "categoria": categoria or "SIN_CATEGORIA",
            "carpeta": str(carpeta) if carpeta else "SIN_CARPETA",
            "advertencias": [],
        }

        if not categoria:
            detalle["advertencias"].append("No se pudo detectar categoria; no se insertaran fotos.")
            return [], detalle

        if categoria not in self.rutas_fotos:
            detalle["advertencias"].append(
                f"No existe carpeta configurada para la categoria {categoria}; no se insertaran fotos."
            )
            return [], detalle

        categoria_carpeta = categoria if categoria in self.rutas_fotos else None
        if categoria != categoria_carpeta:
            detalle["advertencias"].append(
                f"Error de seguridad: categoria detectada {categoria} no coincide "
                f"con la carpeta configurada ({categoria_carpeta or 'SIN_CATEGORIA'}). No se insertaran fotos."
            )
            return [], detalle

        fotos = self.obtener(categoria, cantidad)
        if not fotos:
            detalle["advertencias"].append(
                f"Advertencia: no hay fotos disponibles en la carpeta de {categoria}; "
                "se dejaran los placeholders sin modificar."
            )
        elif categoria == "MINISPLIT" and len(fotos) < len(MINISPLIT_SUBCARPETAS):
            detalle["advertencias"].append(
                "Advertencia: Minisplit requiere una imagen de completo, compresor y partes; "
                "faltan imagenes en una o mas subcarpetas."
            )
        return fotos, detalle


# ================= DOCX =================


def recolectar_placeholders(doc):
    return doc.inline_shapes


def evitar_ruptura_tablas(doc):
    for table in doc.tables:
        for row in table.rows:
            trPr = row._tr.get_or_add_trPr()
            cantSplit = OxmlElement("w:cantSplit")
            cantSplit.set(qn("w:val"), "on")
            trPr.append(cantSplit)


def reemplazar_imagen_shape(doc, shape, ruta):
    try:
        from io import BytesIO

        max_width = shape.width
        max_height = shape.height
        rId = shape._inline.graphic.graphicData.pic.blipFill.blip.embed
        image_part = doc.part.related_parts[rId]
        with PILImage.open(ruta) as img:
            target_ratio = max_width / max_height
            image_ratio = img.width / img.height

            if image_ratio > target_ratio:
                new_width = int(img.height * target_ratio)
                left = (img.width - new_width) // 2
                img = img.crop((left, 0, left + new_width, img.height))
            elif image_ratio < target_ratio:
                new_height = int(img.width / target_ratio)
                top = (img.height - new_height) // 2
                img = img.crop((0, top, img.width, top + new_height))

            buffer = BytesIO()
            formato = "PNG" if img.mode in ("RGBA", "LA", "P") else "JPEG"
            if formato == "JPEG" and img.mode != "RGB":
                img = img.convert("RGB")
            img.save(buffer, format=formato)
            image_part._blob = buffer.getvalue()

        shape.width = max_width
        shape.height = max_height
    except Exception as e:
        print("Error insertando imagen:", e)


def reemplazar_texto(doc, datos):
    reemplazos = {
        "{{FOLIO}}": datos.get("folio", ""),
        "{{REGION}}": datos.get("region", ""),
        "{{UNIDAD}}": datos.get("unidad", ""),
        "{{MARCA}}": datos.get("marca", ""),
        "{{MODELO}}": datos.get("modelo", ""),
        "{{SERIE}}": datos.get("serie", ""),
        "{{INVENTARIO}}": datos.get("inventario", ""),
        "{{EQUIPO}}": datos.get("equipo", ""),
        "{{SERVICIO}}": datos.get("servicio", ""),
        "{{TIPO}}": datos.get("tipo", ""),
    }

    bloques = [doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                bloques.append(cell.paragraphs)

    for grupo in bloques:
        for p in grupo:
            for k, v in reemplazos.items():
                if k in p.text:
                    p.text = p.text.replace(k, str(v))


def marcar_checkbox(doc, servicio):
    variantes = [f"[ ] {servicio}", f"[] {servicio}"]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for v in variantes:
                        if v in p.text:
                            p.text = p.text.replace(v, f"X {servicio}")
                            for run in p.runs:
                                run.font.size = Pt(8.5)


# ================= GENERADOR =================


def generar_reporte(args):
    i, equipo, fotos, template_path, reportes_dir = args
    doc = Document(template_path)
    reemplazar_texto(doc, equipo)
    marcar_checkbox(doc, equipo.get("servicio", ""))

    shapes = recolectar_placeholders(doc)
    for idx, foto in enumerate(fotos or []):
        if idx < len(shapes):
            reemplazar_imagen_shape(doc, shapes[idx], foto)

    evitar_ruptura_tablas(doc)
    nombre = f"{i + 1:04d}_{sanitizar_nombre(equipo.get('folio'))}_{sanitizar_nombre(equipo.get('servicio'))}.docx"
    salida = Path(reportes_dir) / nombre
    doc.save(salida)
    return str(salida)


def preparar_reportes(equipos, rutas_fotos, log=print):
    gestor = GestorFotos(rutas_fotos, log=log)
    trabajos = []
    for i, equipo in enumerate(equipos):
        fotos, detalle = gestor.obtener_para_equipo(equipo, IMAGENES_POR_REPORTE)
        advertencias = detalle.get("advertencias") or []
        if not fotos:
            advertencias.append(
                "Advertencia: el reporte se generara sin fotos porque no hubo fotos disponibles "
                "en la categoria correcta."
            )
        log("---- Seleccion de fotos ----")
        log(f"Folio: {detalle.get('folio', '')}")
        log(f"Equipo: {detalle.get('equipo', '')}")
        log(f"Categoria: {detalle.get('categoria', '')}")
        log(f"Carpeta de fotos: {detalle.get('carpeta', '')}")
        log(f"Fotos seleccionadas: {len(fotos)}")
        log("Advertencias: " + ("; ".join(advertencias) if advertencias else "Ninguna"))
        trabajos.append((i, equipo, fotos))
    return trabajos


def normalizar_rutas_fotos(rutas_fotos):
    normalizadas = {}
    for cat, ruta in (rutas_fotos or {}).items():
        categoria = normalizar_categoria(cat)
        if categoria in categorias_estrictas():
            normalizadas[categoria] = ruta
    return normalizadas


def crear_estructura_salida(unidad, servicio, output_folder=OUTPUT_FOLDER):
    hospital_sanitizado = sanitizar_nombre(unidad)
    servicio_sanitizado = servicio_archivo(servicio)
    hospital_dir = Path(output_folder) / hospital_sanitizado
    reportes_dir = hospital_dir / "reportes_individuales"
    reportes_dir.mkdir(parents=True, exist_ok=True)
    archivo_final = hospital_dir / f"reporte_final_{hospital_sanitizado}_{servicio_sanitizado}.docx"
    log_path = hospital_dir / "log_generacion.txt"
    return hospital_dir, reportes_dir, archivo_final, log_path


def generar_reportes(
    equipos,
    unidad,
    servicio,
    rutas_fotos=None,
    template_path=TEMPLATE_PATH,
    output_folder=OUTPUT_FOLDER,
    usar_multiproceso=True,
    log=print,
):
    equipos = list(equipos)
    if rutas_fotos is None:
        rutas_fotos = RUTAS_FOTOS
    else:
        rutas_fotos = normalizar_rutas_fotos(rutas_fotos)
    hospital_dir, reportes_dir, archivo_final, log_path = crear_estructura_salida(
        unidad, servicio, output_folder
    )
    if log_path.exists():
        log_path.unlink()

    def log_dual(mensaje):
        log(mensaje)
        escribir_log(log_path, mensaje)

    log_dual(f"Hospital seleccionado: {unidad}")
    log_dual(f"Tipo de servicio: {servicio}")
    preventivos, correctivos = contar_servicios(equipos)
    log_dual(f"Total de equipos encontrados: {len(equipos)}")
    log_dual(f"Total preventivos: {preventivos}")
    log_dual(f"Total correctivos: {correctivos}")
    log_dual("Rutas de fotos activas:")
    for categoria in categorias_estrictas():
        ruta = rutas_fotos.get(categoria, "SIN_RUTA") if isinstance(rutas_fotos, dict) else "SIN_RUTA"
        log_dual(f"{categoria} -> {ruta}")

    if not equipos:
        log_dual("No se encontraron equipos para los filtros seleccionados.")
        return None, 0, str(log_path)

    log_dual("Generando reportes individuales...")
    trabajos = preparar_reportes(equipos, rutas_fotos, log=log_dual)
    args = [
        (i, equipo, fotos, str(template_path), str(reportes_dir))
        for i, equipo, fotos in trabajos
    ]

    if usar_multiproceso and len(args) > 1:
        workers = min(max(1, multiprocessing.cpu_count() - 1), len(args))
        with ProcessPoolExecutor(workers) as executor:
            archivos = list(executor.map(generar_reporte, args))
    else:
        archivos = [generar_reporte(arg) for arg in args]

    archivos = [archivo for archivo in archivos if archivo]
    if not archivos:
        log_dual("No se pudo generar ningun documento individual.")
        return None, 0, str(log_path)

    log_dual("Uniendo documentos con docxcompose...")
    master = Document(archivos[0])
    composer = Composer(master)
    for archivo in archivos[1:]:
        composer.append(Document(archivo))
    composer.save(archivo_final)

    log_dual(f"Reportes generados: {len(archivos)}")
    log_dual(f"Archivo final: {archivo_final}")
    log_dual(f"Carpeta de salida: {hospital_dir}")
    return str(archivo_final), len(archivos), str(log_path)


# ================= INTERFAZ VISUAL =================


class GeneradorApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Generador de reportes de servicio")
        self.root.geometry("950x700")
        self.root.minsize(820, 620)

        self.db_resumen = inicializar_db()

        self.template_var = tk.StringVar(value=str(TEMPLATE_PATH))
        self.output_var = tk.StringVar(value=str(OUTPUT_FOLDER))
        self.busqueda_var = tk.StringVar()
        self.unidad_var = tk.StringVar()
        self.servicio_var = tk.StringVar(value="Ambos")
        self.resumen_var = tk.StringVar(value="Busca y selecciona una unidad medica.")
        self.ruta_vars = {
            categoria: tk.StringVar(value=str(ruta))
            for categoria, ruta in RUTAS_FOTOS.items()
        }
        self.unidades_actuales = []
        self._crear_vista()
        self._mostrar_resumen_base()

    def _crear_vista(self):
        contenedor = ttk.Frame(self.root, padding=16)
        contenedor.pack(fill="both", expand=True)
        contenedor.columnconfigure(1, weight=1)

        ttk.Label(contenedor, text="Plantilla Word").grid(row=0, column=0, sticky="w", pady=4)
        ttk.Entry(contenedor, textvariable=self.template_var).grid(row=0, column=1, sticky="ew", padx=8, pady=4)
        ttk.Button(contenedor, text="Elegir", command=self._elegir_plantilla).grid(row=0, column=2, pady=4)

        ttk.Label(contenedor, text="Carpeta de salida").grid(row=1, column=0, sticky="w", pady=4)
        ttk.Entry(contenedor, textvariable=self.output_var).grid(row=1, column=1, sticky="ew", padx=8, pady=4)
        ttk.Button(contenedor, text="Elegir", command=self._elegir_salida).grid(row=1, column=2, pady=4)

        ttk.Label(contenedor, text="Buscar hospital").grid(row=2, column=0, sticky="w", pady=4)
        busqueda = ttk.Entry(contenedor, textvariable=self.busqueda_var)
        busqueda.grid(row=2, column=1, sticky="ew", padx=8, pady=4)
        busqueda.bind("<KeyRelease>", lambda _event: self._actualizar_unidades())
        ttk.Button(contenedor, text="Buscar", command=self._actualizar_unidades).grid(row=2, column=2, pady=4)

        ttk.Label(contenedor, text="Unidad medica").grid(row=3, column=0, sticky="w", pady=4)
        self.unidad_combo = ttk.Combobox(contenedor, textvariable=self.unidad_var, state="readonly")
        self.unidad_combo.grid(row=3, column=1, sticky="ew", padx=8, pady=4)
        self.unidad_combo.bind("<<ComboboxSelected>>", lambda _event: self._actualizar_resumen())

        ttk.Label(contenedor, text="Servicio").grid(row=4, column=0, sticky="w", pady=4)
        servicio_combo = ttk.Combobox(
            contenedor,
            textvariable=self.servicio_var,
            values=SERVICIOS,
            state="readonly",
            width=18,
        )
        servicio_combo.grid(row=4, column=1, sticky="w", padx=8, pady=4)
        servicio_combo.bind("<<ComboboxSelected>>", lambda _event: self._actualizar_resumen())

        resumen = ttk.Label(contenedor, textvariable=self.resumen_var, justify="left")
        resumen.grid(row=5, column=0, columnspan=3, sticky="ew", pady=(8, 4))

        marco_fotos = ttk.LabelFrame(contenedor, text="Carpetas de fotos", padding=10)
        marco_fotos.grid(row=6, column=0, columnspan=3, sticky="nsew", pady=12)
        marco_fotos.columnconfigure(1, weight=1)

        for fila, (categoria, variable) in enumerate(self.ruta_vars.items()):
            ttk.Label(marco_fotos, text=categoria).grid(row=fila, column=0, sticky="w", pady=3)
            ttk.Entry(marco_fotos, textvariable=variable).grid(row=fila, column=1, sticky="ew", padx=8, pady=3)
            ttk.Button(
                marco_fotos,
                text="Elegir",
                command=lambda var=variable: self._elegir_carpeta(var),
            ).grid(row=fila, column=2, pady=3)

        acciones = ttk.Frame(contenedor)
        acciones.grid(row=7, column=0, columnspan=3, sticky="ew", pady=8)
        acciones.columnconfigure(0, weight=1)
        ttk.Button(acciones, text="Reconstruir base local", command=self._reconstruir_base).grid(row=0, column=0, sticky="w")
        self.boton_generar = ttk.Button(acciones, text="Generar reportes", command=self._generar)
        self.boton_generar.grid(row=0, column=1, sticky="e")

        self.log_text = tk.Text(contenedor, height=10, wrap="word")
        self.log_text.grid(row=8, column=0, columnspan=3, sticky="nsew", pady=(8, 0))
        contenedor.rowconfigure(8, weight=1)

        self._actualizar_unidades()

    def _mostrar_resumen_base(self):
        resumen = self.db_resumen or inicializar_db()
        if resumen.get("advertencia"):
            self._log(resumen["advertencia"])
        self._log(f"Base de datos cargada: {resumen.get('source_file')}")
        self._log(f"Total de registros: {resumen.get('total_registros')}")
        self._log(f"Total de unidades: {resumen.get('total_unidades')}")

    def _reconstruir_base(self):
        if not messagebox.askyesno(
            "Reconstruir base local",
            "Esto borrara y recreara equipos.db usando la base corregida. Deseas continuar?",
        ):
            return
        try:
            self.db_resumen = inicializar_db(force_rebuild=True)
            self._actualizar_unidades()
            self._mostrar_resumen_base()
            messagebox.showinfo("Base reconstruida", "La base local se reconstruyo correctamente.")
        except Exception as exc:
            messagebox.showerror("Error", str(exc))

    def _elegir_plantilla(self):
        ruta = filedialog.askopenfilename(
            title="Elegir plantilla",
            filetypes=[("Word", "*.docx"), ("Todos", "*.*")],
            initialdir=str(BASE_DIR),
        )
        if ruta:
            self.template_var.set(ruta)

    def _elegir_salida(self):
        ruta = filedialog.askdirectory(title="Elegir carpeta de salida", initialdir=str(BASE_DIR))
        if ruta:
            self.output_var.set(ruta)

    def _elegir_carpeta(self, variable):
        ruta = filedialog.askdirectory(title="Elegir carpeta de fotos", initialdir=str(BASE_DIR))
        if ruta:
            variable.set(ruta)

    def _actualizar_unidades(self):
        self.unidades_actuales = buscar_unidades(self.busqueda_var.get())
        self.unidad_combo["values"] = self.unidades_actuales
        if self.unidades_actuales:
            if self.unidad_var.get() not in self.unidades_actuales:
                self.unidad_var.set(self.unidades_actuales[0])
        else:
            self.unidad_var.set("")
        self._actualizar_resumen()

    def _equipos_seleccionados(self):
        unidad = self.unidad_var.get()
        if not unidad:
            return []
        return obtener_equipos(unidad, self.servicio_var.get())

    def _actualizar_resumen(self):
        unidad = self.unidad_var.get()
        servicio = self.servicio_var.get()
        equipos = self._equipos_seleccionados()
        preventivos, correctivos = contar_servicios(equipos)
        self.resumen_var.set(
            "Hospital seleccionado: {unidad}\n"
            "Tipo de servicio: {servicio}\n"
            "Total de equipos encontrados: {total}\n"
            "Total preventivos: {preventivos}\n"
            "Total correctivos: {correctivos}".format(
                unidad=unidad or "Sin seleccion",
                servicio=servicio,
                total=len(equipos),
                preventivos=preventivos,
                correctivos=correctivos,
            )
        )

    def _log(self, mensaje):
        if threading.current_thread() is not threading.main_thread():
            self.root.after(0, self._log, mensaje)
            return
        self.log_text.insert("end", str(mensaje) + "\n")
        self.log_text.see("end")
        self.root.update_idletasks()

    def _generar(self):
        unidad = self.unidad_var.get()
        servicio = self.servicio_var.get()
        equipos = self._equipos_seleccionados()

        if not unidad:
            messagebox.showerror("Unidad requerida", "Busca y selecciona una unidad medica.")
            return
        if not equipos:
            messagebox.showwarning("Sin equipos", "No hay equipos para la unidad y servicio seleccionados.")
            return

        template_path = self.template_var.get().strip()
        output_folder = self.output_var.get().strip()
        if not os.path.exists(template_path):
            messagebox.showerror("Plantilla no encontrada", "Selecciona una plantilla Word valida.")
            return

        rutas_fotos = {categoria: var.get().strip() for categoria, var in self.ruta_vars.items()}

        self.boton_generar.config(state="disabled")
        self.log_text.delete("1.0", "end")
        self._actualizar_resumen()

        def trabajo():
            try:
                resultado, total, log_path = generar_reportes(
                    equipos=equipos,
                    unidad=unidad,
                    servicio=servicio,
                    rutas_fotos=rutas_fotos,
                    template_path=Path(template_path),
                    output_folder=Path(output_folder),
                    usar_multiproceso=True,
                    log=self._log,
                )
                if resultado:
                    self.root.after(0, messagebox.showinfo, "Listo", f"Se generaron {total} reportes.\n{resultado}\nLog: {log_path}")
                else:
                    self.root.after(0, messagebox.showwarning, "Sin reportes", f"No se pudo generar ningun reporte.\nLog: {log_path}")
            except Exception as exc:
                self._log(f"Error: {exc}")
                self.root.after(0, messagebox.showerror, "Error", str(exc))
            finally:
                self.root.after(0, self.boton_generar.config, {"state": "normal"})

        threading.Thread(target=trabajo, daemon=True).start()


def lanzar_interfaz():
    root = tk.Tk()
    GeneradorApp(root)
    root.mainloop()


def main():
    inicializar_db()
    print("Base de datos lista:", DB_PATH)


if __name__ == "__main__":
    multiprocessing.freeze_support()
    lanzar_interfaz()
